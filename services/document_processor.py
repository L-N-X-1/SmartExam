# services/document_processor.py
"""
Extraction de texte depuis PDF, DOCX et TXT
"""

# Import flexible pour PDF
from docx import Document
import os
try:
    import PyPDF2
except ImportError:
    import pypdf as PyPDF2



def extract_text_from_pdf(file_path):
    """
    Extrait le texte d'un fichier PDF
    INPUT: file_path (str) - chemin vers le PDF
    OUTPUT: text (str) - texte complet extrait
    """
    text = ""
    try:
        with open(file_path, 'rb') as file:
            pdf_reader = PyPDF2.PdfReader(file)
            num_pages = len(pdf_reader.pages)
            print(f"   📄 PDF : {num_pages} page(s)")
            
            for i, page in enumerate(pdf_reader.pages, 1):
                page_text = page.extract_text()
                if page_text:
                    text += page_text + "\n"
                    
    except Exception as e:
        print(f"   ❌ Erreur lecture PDF: {e}")
        
    return text.strip()

def extract_text_from_docx(file_path):
    """
    Extrait le texte d'un fichier DOCX
    INPUT: file_path (str) - chemin vers le DOCX
    OUTPUT: text (str) - texte complet extrait
    """
    text = ""
    try:
        doc = Document(file_path)
        num_paragraphs = len(doc.paragraphs)
        print(f"   📄 DOCX : {num_paragraphs} paragraphe(s)")
        
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text += paragraph.text + "\n"
                
    except Exception as e:
        print(f"   ❌ Erreur lecture DOCX: {e}")
        
    return text.strip()

def extract_text_from_txt(file_path):
    """
    Extrait le texte d'un fichier TXT
    INPUT: file_path (str) - chemin vers le TXT
    OUTPUT: text (str) - texte complet extrait
    """
    try:
        with open(file_path, 'r', encoding='utf-8') as f:
            text = f.read()
        print(f"   📄 TXT : {len(text)} caractères")
        return text.strip()
    except UnicodeDecodeError:
        # Essaie avec un autre encodage
        try:
            with open(file_path, 'r', encoding='latin-1') as f:
                text = f.read()
            print(f"   📄 TXT (latin-1) : {len(text)} caractères")
            return text.strip()
        except Exception as e:
            print(f"   ❌ Erreur lecture TXT: {e}")
            return ""
    except Exception as e:
        print(f"   ❌ Erreur lecture TXT: {e}")
        return ""

def process_document(file_path):
    """
    Détecte le type de fichier et extrait le texte
    INPUT: file_path (str) - chemin vers le fichier
    OUTPUT: text (str) - texte extrait
    
    Formats supportés: PDF, DOCX, TXT
    """
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"❌ Fichier introuvable: {file_path}")
    
    file_path_lower = file_path.lower()
    file_name = os.path.basename(file_path)
    
    print(f"\n📖 Traitement : {file_name}")
    
    if file_path_lower.endswith('.pdf'):
        return extract_text_from_pdf(file_path)
    elif file_path_lower.endswith('.docx'):
        return extract_text_from_docx(file_path)
    elif file_path_lower.endswith('.txt'):
        return extract_text_from_txt(file_path)
    else:
        extension = os.path.splitext(file_path)[1]
        raise ValueError(
            f"❌ Format non supporté: {extension}\n"
            f"   Formats acceptés: .pdf, .docx, .txt"
        )